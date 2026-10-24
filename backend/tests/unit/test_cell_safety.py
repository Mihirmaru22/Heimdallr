import pytest
from openpyxl import load_workbook

from app.export.cell_safety import EXCEL_MAX_CELL, cell_value, safe_sheet_name, safe_text
from app.export.xlsx import ExportOptions, ExportRecord, build_workbook
from app.spec.models import ExtractionSpec

PAYLOADS = ["=1+1", "+cmd|' /C calc'!A0", "-2+3", "@SUM(A1:A2)", "\t=1", "\r=1", "＝1+1", " =HYPERLINK(\"x\")"]


@pytest.mark.parametrize("payload", PAYLOADS)
def test_formula_prefixes_neutralized(payload):
    assert safe_text(payload).startswith("'")


def test_plain_and_numeric_text_untouched():
    assert safe_text("Acme Ltd") == "Acme Ltd"
    assert safe_text("-42") == "-42"
    assert cell_value(3.5) == 3.5
    assert cell_value(["=a", "b"]).startswith("'")


def test_long_values_truncated_to_excel_limit():
    assert len(safe_text("x" * 40_000)) <= EXCEL_MAX_CELL


def test_sheet_names_are_valid_and_unique():
    used: set[str] = set()
    assert safe_sheet_name("a/b:c", used) == "a b c"
    assert safe_sheet_name("A B C", used) != "A B C"
    assert len(safe_sheet_name("x" * 50, used)) <= 31


def test_workbook_contains_no_formulas(tmp_path):
    spec = ExtractionSpec.model_validate({"fields": [
        {"name": "company_name", "type": "string", "semantic": "company_name", "required": True},
        {"name": "description", "type": "text"},
        {"name": "website_url", "type": "url"},
    ]})
    records = [ExportRecord(str(i), {"company_name": p, "description": p, "website_url": "https://x.example/"},
                            "valid", 1.0, ["https://x.example/"],
                            [{"field_name": "description", "value": p, "value_status": "found", "method": "meta",
                              "source_url": "https://x.example/", "evidence": p, "confidence": 0.8, "error": None}])
               for i, p in enumerate(PAYLOADS)]
    summary = {"Request": "=EVIL()", "Rows": len(records)}
    result = build_workbook(tmp_path / "out.xlsx", spec, records, summary, ExportOptions())
    assert result.rows == len(PAYLOADS)
    wb = load_workbook(result.path)
    assert {"Data", "Sources", "Run Summary"} <= set(wb.sheetnames)
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                assert cell.data_type != "f", f"formula in {ws.title}!{cell.coordinate}: {cell.value!r}"
                if isinstance(cell.value, str):
                    assert not cell.value.lstrip(" ").startswith(("=", "+", "@", "＝")), cell.value


def test_phone_numbers_are_not_prefixed():
    assert safe_text("+18663001261") == "+18663001261"
    assert safe_text("+1 (866) 300-1261") == "+1 (866) 300-1261"
    assert safe_text("+1+cmd|' /C calc'!A0").startswith("'")
