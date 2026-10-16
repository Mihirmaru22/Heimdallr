"""Test suite for Context-Aware Relevance, Narrative Synthesis, and Dual Exporters.

Verifies:
1. Contextual relevance does NOT act as a crude domain blacklist:
   - Legitimate infrastructure for Google/YouTube is accepted when investigating them or CDN/streaming.
   - Legitimate target infrastructure (e.g. paramsoftware.com) is accepted for investigated target.
   - Incidental hosting platform CDNs (e.g. googlevideo.com on an unrelated subject's YouTube video) are rejected.
2. Narrative intelligence synthesizes human-readable prose rather than raw relational triples.
3. Excel exporter outputs executive brief and categorized findings in Tab 1, and filtered entities in Tab 3.
4. Markdown exporter generates structured, human-readable intelligence briefs.
"""

from __future__ import annotations

from pathlib import Path
import tempfile
import openpyxl
import pytest

from argus.investigation.relevance import (
    clean_entity_name,
    is_contextually_relevant_host,
)
from argus.export.narrative import generate_narrative_intelligence
from argus.export.excel_exporter import OSINTExcelExporter
from argus.export.markdown_exporter import OSINTMarkdownExporter


class TestContextAwareRelevance:
    """Verifies that infrastructure filtering is context and relevance-driven, never a crude blacklist."""

    def test_target_infrastructure_accepted(self):
        """Target's own domain and endpoints must always be accepted."""
        target = "Param Software"
        host = "paramsoftware.com"
        assert is_contextually_relevant_host(
            host,
            doc_url="https://paramsoftware.com",
            target_objective="Find everything about Param Software",
            target_entities=[target],
        ) is True

    def test_google_youtube_accepted_when_investigating_google_or_youtube(self):
        """Google and YouTube domains must NOT be blacklisted when they are the subject of investigation."""
        # 1. Direct target is Google
        assert is_contextually_relevant_host(
            "googlevideo.com",
            doc_url="https://youtube.com",
            target_objective="Investigate Google CDN architecture",
            target_entities=["Google"],
        ) is True

        # 2. Target is YouTube
        assert is_contextually_relevant_host(
            "rr1---sn-ab5sznzl.googlevideo.com",
            doc_url="https://youtube.com/watch?v=xyz",
            target_objective="Audit YouTube video streaming infrastructure",
            target_entities=["YouTube"],
        ) is True

        # 3. Investigation objective explicitly inquires into streaming/CDN/network infrastructure
        assert is_contextually_relevant_host(
            "googlevideo.com",
            doc_url="https://youtube.com/watch?v=xyz",
            target_objective="Investigate global streaming and CDN video delivery topologies",
            target_entities=["Video Delivery"],
        ) is True

    def test_incidental_platform_cdn_rejected_for_unrelated_target(self):
        """When investigating an unrelated third-party, platform media CDNs are recognized as incidental."""
        # Target is an executive / software company
        target = "Ronak Raiyani"
        obj = "Find everything about Ronak Raiyani Param Software"

        # Host is a video streaming server embedded in YouTube HTML
        host = "rr1---sn-4g5ednss.googlevideo.com"
        doc_url = "https://www.youtube.com/watch?v=nN43Y27bWss"

        assert is_contextually_relevant_host(
            host,
            doc_url=doc_url,
            target_objective=obj,
            target_entities=[target, "Param Software"],
        ) is False

    def test_entity_name_cleaning(self):
        """Interview prefixes and boilerplate titles must be stripped cleanly."""
        assert clean_entity_name("In Conversation with Mr. Ronak Raiyani") == "Ronak Raiyani"
        assert clean_entity_name("In Conversation with Dr Jane Doe") == "Jane Doe"
        assert clean_entity_name("Meet the Maker | John Smith") == "John Smith"
        assert clean_entity_name("Know Our IACC Member: Param Software") == "Param Software"


class TestNarrativeAndExporters:
    """Verifies human-readable narrative intelligence generation and dual exports."""

    @pytest.fixture
    def mock_osint_data(self):
        return {
            "investigation_id": "inv_test_001",
            "target": "Ronak Raiyani",
            "objective": "Find everything about Ronak Raiyani Param Software",
            "confidence": 0.94,
            "is_stopped": True,
            "stop_reason": "Information Gain Saturation",
            "evidence": [
                {
                    "title": "Ronak Raiyani - Founder and CTO at Param Software",
                    "url": "https://www.linkedin.com/in/ronakraiyani",
                    "snippet": "Ronak Raiyani is Founder & CTO at Param Software. President of Rajkot Information Technology Association (RITA). Certified Scrum Master. Headquartered in Rajkot, Gujarat, India. Contact: info@paramsoftware.com",
                },
                {
                    "title": "In Conversation with Mr. Ronak Raiyani - YouTube",
                    "url": "https://www.youtube.com/watch?v=123",
                    "snippet": "Watch our interview with Ronak Raiyani of Param Software discussing cloud scalability and enterprise engineering in Gujarat.",
                },
            ],
            "claims": [
                {
                    "subject": "Ronak Raiyani",
                    "predicate": "founder_and_cto",
                    "object": "Param Software",
                    "status": "SUPPORTED",
                    "confidence": 0.95,
                    "citation": "https://www.linkedin.com/in/ronakraiyani",
                },
                {
                    "subject": "In Conversation with Mr. Ronak Raiyani",
                    "predicate": "infrastructure",
                    "object": "rr1---sn-4g5ednss.googlevideo.com",
                    "status": "SUPPORTED",
                    "confidence": 0.85,
                    "citation": "https://www.youtube.com/watch?v=123",
                },
            ],
            "entities": {
                "Ronak Raiyani": {"entity_type": "INDIVIDUAL", "role": "Founder & CTO"},
                "Param Software": {"entity_type": "ORGANIZATION", "role": "Operating Enterprise"},
                "rr1---sn-4g5ednss.googlevideo.com": {"entity_type": "INFRASTRUCTURE_ENDPOINT"},
            },
            "research_actions": [{"action": "SEARCH"}, {"action": "FETCH"}],
        }

    def test_narrative_intelligence_synthesis(self, mock_osint_data):
        """Synthesized narrative must contain prose executive overview and categorized findings."""
        synth = generate_narrative_intelligence(mock_osint_data)

        # Primary name resolution
        assert synth["primary_name"] == "Ronak Raiyani"

        # Executive narrative prose
        narrative = synth["executive_narrative"]
        assert "Ronak Raiyani" in narrative
        assert "Param Software" in narrative
        assert "Founder" in narrative or "CTO" in narrative

        # Categorized findings
        cat = synth["categorized_findings"]
        assert "Executive & Professional Identity" in cat
        assert any("CTO" in f["fact"] or "Founder" in f["fact"] for f in cat["Executive & Professional Identity"])
        assert any("Rajkot" in f["fact"] for f in cat["Corporate & Organizational Footprint"])

        # Curated entities must NOT include the YouTube CDN host
        curated_names = [e["name"] for e in synth["curated_entities"]]
        assert "Ronak Raiyani" in curated_names
        assert "Param Software" in curated_names
        assert "rr1---sn-4g5ednss.googlevideo.com" not in curated_names

    def test_excel_export_human_readable(self, mock_osint_data):
        """Excel exporter must write executive brief and clean entities without crude triples in summary."""
        with tempfile.TemporaryDirectory() as tmpdir:
            exporter = OSINTExcelExporter(export_dir=tmpdir)
            out_file = exporter.export_investigation(mock_osint_data)

            assert out_file.exists()
            wb = openpyxl.load_workbook(out_file)

            # Tab 1: Executive Summary
            ws_summary = wb["Executive Summary"]
            cell_texts = [str(ws_summary.cell(row=r, column=1).value or "") for r in range(1, 30)]
            full_summary_text = "\n".join(cell_texts)

            assert "EXECUTIVE INTELLIGENCE BRIEF" in full_summary_text
            assert "KEY CORROBORATED FINDINGS" in full_summary_text
            # Raw triple arrow format must NOT be used
            assert "• Subj -> pred: obj" not in full_summary_text
            assert "-> Founder And Cto:" not in full_summary_text

            # Tab 3: Entities Dossier must NOT contain the incidental YouTube CDN
            ws_entities = wb["Entities Dossier"]
            entity_names = [str(ws_entities.cell(row=r, column=1).value or "") for r in range(2, 20)]
            assert "Ronak Raiyani" in entity_names
            assert "rr1---sn-4g5ednss.googlevideo.com" not in entity_names

    def test_markdown_export(self, mock_osint_data):
        """Markdown exporter must produce a clean, readable intelligence brief."""
        with tempfile.TemporaryDirectory() as tmpdir:
            exporter = OSINTMarkdownExporter(export_dir=tmpdir)
            out_file = exporter.export_investigation(mock_osint_data)

            assert out_file.exists()
            content = out_file.read_text(encoding="utf-8")

            assert "# HEIMDALLR V4 — OSINT INTELLIGENCE DOSSIER" in content
            assert "## Executive Intelligence Brief" in content
            assert "## Key Corroborated Findings" in content
            assert "## Identified Entities & Operational Profiles" in content
            assert "Ronak Raiyani" in content
            assert "Param Software" in content
            assert "rr1---sn-4g5ednss.googlevideo.com" not in content

    def test_technographic_deployments_extraction_and_epistemic_tiers(self):
        """Verify extraction of company deployments, attributes, and 3-tier epistemic provenance."""
        from argus.investigation.technographics import (
            extract_technographic_deployments,
            classify_technographic_source,
        )

        theirstack_text = (
            "API Export Company Country Industry Employees Revenue Technologies "
            "Northrop Grumman United States Defense and Space Manufacturing 95k $39B Siemens Teamcenter "
            "Honeywell United States Appliances, Electrical, and Electronics Manufacturing 95k $36B Siemens Teamcenter "
            "Ford Motor Company United States Motor Vehicle Manufacturing 177k $176B Siemens Teamcenter "
            "Hermeus United States Aviation and Aerospace Component Manufacturing 140 Siemens Teamcenter"
        )
        provider_url = "https://theirstack.com/en/technology/siemens-teamcenter/us"

        findings = extract_technographic_deployments(
            theirstack_text,
            doc_url=provider_url,
            target_query="Search Companies who are using TEAMCENTER in usa",
        )
        assert len(findings) == 4
        comp_map = {f.company: f for f in findings}

        # Check Northrop Grumman
        ng = comp_map["Northrop Grumman"]
        assert ng.country == "United States"
        assert ng.industry == "Defense and Space Manufacturing"
        assert ng.employees == "95k"
        assert ng.revenue == "$39B"
        assert ng.software == "Siemens Teamcenter"
        assert ng.status == "TECHNOGRAPHIC_PROVIDER_ASSERTION"
        assert ng.source_url == provider_url

        # Check Hermeus (no revenue reported)
        hermeus = comp_map["Hermeus"]
        assert hermeus.country == "United States"
        assert hermeus.industry == "Aviation and Aerospace Component Manufacturing"
        assert hermeus.employees == "140"
        assert hermeus.revenue == ""

        # Verify Vendor provenance classification
        v_status, v_conf, v_type, _ = classify_technographic_source(
            "https://sw.siemens.com/en-US/case-studies/ford-motor/", "Siemens Teamcenter"
        )
        assert v_status == "VENDOR_CONFIRMED"
        # A publisher's class is not a probability; no confidence is fabricated.
        assert v_conf is None

        # Verify Technographic Provider provenance classification
        p_status, p_conf, p_type, _ = classify_technographic_source("https://theirstack.com/en/technology/siemens-teamcenter/us")
        assert p_status == "TECHNOGRAPHIC_PROVIDER_ASSERTION"

    def test_technographic_narrative_intelligence_sample_labeling(self):
        """Verify that technographic investigation produces a sample-labeled executive summary."""
        tech_data = {
            "investigation_id": "inv_tech_001",
            "target": "Search Companies who are using TEAMCENTER in usa",
            "objective": "Search Companies who are using TEAMCENTER in usa",
            "confidence": 0.92,
            "is_stopped": True,
            "stop_reason": "Information Gain Saturation",
            "evidence": [
                {
                    "title": "Companies using Siemens Teamcenter in United States",
                    "url": "https://theirstack.com/en/technology/siemens-teamcenter/us",
                    "snippet": "Northrop Grumman United States Defense and Space Manufacturing 95k $39B Siemens Teamcenter",
                }
            ],
            "claims": [
                {
                    "subject": "Northrop Grumman",
                    "predicate": "USES_PLM_SOFTWARE",
                    "object": "Siemens Teamcenter",
                    "status": "TECHNOGRAPHIC_PROVIDER_ASSERTION",
                    "confidence": 0.88,
                    "source_url": "https://theirstack.com/en/technology/siemens-teamcenter/us",
                    "country": "United States",
                    "industry": "Defense and Space Manufacturing",
                    "employees": "95k",
                    "revenue": "$39B",
                },
                {
                    "subject": "Lockheed Martin",
                    "predicate": "USES_PLM_SOFTWARE",
                    "object": "Siemens Teamcenter",
                    "status": "VENDOR_CONFIRMED",
                    "confidence": 0.96,
                    "source_url": "https://news.siemens.com/press-release/lockheed-martin-teamcenter",
                    "country": "United States",
                    "industry": "Defense and Space Manufacturing",
                    "employees": "122k",
                    "revenue": "$69B",
                },
            ],
            "entities": {
                "Northrop Grumman": {
                    "entity_type": "ORGANIZATION",
                    "metadata": {"country": "United States", "industry": "Defense and Space Manufacturing"},
                },
                "Lockheed Martin": {
                    "entity_type": "ORGANIZATION",
                    "metadata": {"country": "United States", "industry": "Defense and Space Manufacturing"},
                },
            },
        }

        synth = generate_narrative_intelligence(tech_data)

        # Primary name must reflect technographic deployment target rather than raw query string
        assert synth["primary_name"] == "Siemens Teamcenter Enterprise Deployments (United States)"

        # Executive narrative must explicitly label findings as a sample
        narrative = synth["executive_narrative"]
        assert "sample" in narrative.lower()
        assert "Siemens Teamcenter" in narrative
        assert "United States" in narrative
        assert "three-tier epistemic provenance" in narrative or "VENDOR_CONFIRMED" in narrative

        # Categorized findings must include the sample size in header
        cat = synth["categorized_findings"]
        cat_key = next(k for k in cat if "Reported Enterprise Deployments" in k)
        assert "[Sample: 2]" in cat_key
        assert any("Northrop Grumman" in f["fact"] and "$39B" in f["fact"] for f in cat[cat_key])
        assert any("Lockheed Martin" in f["fact"] and "VENDOR_CONFIRMED" == f["status"] for f in cat[cat_key])

        # Curated entities must register companies as enterprise deployers
        curated = synth["curated_entities"]
        comp_names = [e["name"] for e in curated]
        assert "Northrop Grumman" in comp_names
        assert "Lockheed Martin" in comp_names
        assert all(e["type"] == "ORGANIZATION" for e in curated)

    def test_technographic_exporters_and_excel_styles(self):
        """Verify that Excel and Markdown exporters render technographic findings cleanly."""
        tech_data = {
            "investigation_id": "inv_tech_002",
            "target": "Search Companies who are using TEAMCENTER in usa",
            "objective": "Search Companies who are using TEAMCENTER in usa",
            "confidence": 0.92,
            "is_stopped": True,
            "stop_reason": "Information Gain Saturation",
            "evidence": [
                {
                    "title": "Siemens Teamcenter US Customers",
                    "url": "https://theirstack.com/en/technology/siemens-teamcenter/us",
                    "snippet": "Ford Motor Company United States Motor Vehicle Manufacturing 177k $176B Siemens Teamcenter",
                }
            ],
            "claims": [
                {
                    "subject": "Ford Motor Company",
                    "predicate": "USES_PLM_SOFTWARE",
                    "object": "Siemens Teamcenter",
                    "status": "TECHNOGRAPHIC_PROVIDER_ASSERTION",
                    "confidence": 0.88,
                    "source_url": "https://theirstack.com/en/technology/siemens-teamcenter/us",
                    "country": "United States",
                    "industry": "Motor Vehicle Manufacturing",
                    "employees": "177k",
                    "revenue": "$176B",
                }
            ],
            "entities": {
                "Ford Motor Company": {
                    "entity_type": "ORGANIZATION",
                    "metadata": {"country": "United States", "industry": "Motor Vehicle Manufacturing"},
                }
            },
            "research_actions": [{"action": "SEARCH"}],
        }

        with tempfile.TemporaryDirectory() as tmpdir:
            excel_exporter = OSINTExcelExporter(export_dir=tmpdir)
            md_exporter = OSINTMarkdownExporter(export_dir=tmpdir)

            excel_path = excel_exporter.export_investigation(tech_data)
            md_path = md_exporter.export_investigation(tech_data)

            # Check Excel content
            wb = openpyxl.load_workbook(excel_path)
            ws_claims = wb["Verified Claims"]
            assert ws_claims.cell(row=2, column=2).value == "Ford Motor Company"
            assert ws_claims.cell(row=2, column=5).value == "TECHNOGRAPHIC_PROVIDER_ASSERTION"
            assert "theirstack.com" in str(ws_claims.cell(row=2, column=7).value)

            # Check Markdown content
            md_content = md_path.read_text(encoding="utf-8")
            assert "siemens teamcenter enterprise deployments (united states)" in md_content.lower()
            assert "Ford Motor Company" in md_content
            assert "TECHNOGRAPHIC_PROVIDER_ASSERTION" in md_content
            assert "sample" in md_content.lower()
