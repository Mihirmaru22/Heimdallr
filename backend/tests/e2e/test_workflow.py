"""The whole workflow through the HTTP API: request -> schema -> job -> records -> Excel download."""

import io

from openpyxl import load_workbook


def test_request_to_excel(make_user, crawl_env):
    alice = make_user("alice@example.com", org="Alice Co")
    project = alice.post(alice.org("/projects"), json={"name": "Cyber vendors"}).json()

    url = crawl_env["url"]
    request = ("Collect the company name, website, contact email and phone of these cybersecurity companies: "
               f"{url('acme-security.test')} {url('globex.test')}")
    r = alice.post(alice.org(f"/projects/{project['id']}/schemas/generate"), json={"request": request})
    assert r.status_code == 201, r.text
    schema = r.json()
    names = [f["name"] for f in schema["spec"]["fields"]]
    assert {"company_name", "website_url", "contact_email", "phone"} <= set(names)
    assert schema["spec"]["discovery"]["providers"] == ["seeds"]

    preview = alice.get(alice.org(f"/schemas/{schema['id']}/preview")).json()
    assert all(u["allowed"] for u in preview["urls"]) and preview["estimated_llm_calls"] == 0

    r = alice.post(alice.org("/jobs"), json={"schema_id": schema["id"]}, headers={"Idempotency-Key": "run-1"})
    assert r.status_code == 201, r.text
    job = r.json()
    # Same key again returns the same job instead of starting another crawl.
    again = alice.post(alice.org("/jobs"), json={"schema_id": schema["id"]}, headers={"Idempotency-Key": "run-1"})
    assert again.status_code == 200 and again.json()["id"] == job["id"]

    progress = alice.get(alice.org(f"/jobs/{job['id']}/progress")).json()
    assert progress["terminal"] and progress["state"] in ("completed", "partially_completed"), progress
    events = alice.get(alice.org(f"/jobs/{job['id']}/events")).json()
    assert [e["to_state"] for e in events if e["type"] == "transition"][:3] == ["created", "validating", "planning"]

    recs = alice.get(alice.org(f"/jobs/{job['id']}/records"), params={"sort": "quality"}).json()
    by_name = {r["data"]["company_name"]: r for r in recs["items"]}
    assert "Acme Security Ltd" in by_name and "Globex Cyber Defense" in by_name

    detail = alice.get(alice.org(f"/records/{by_name['Acme Security Ltd']['id']}")).json()
    assert any(f["field_name"] == "company_name" and f["selected"] and f["source_url"] for f in detail["fields"])

    # A manual correction is validated and kept.
    bad = alice.patch(alice.org(f"/records/{detail['id']}"), json={"data": {"contact_email": "nope"}})
    assert bad.status_code == 422
    ok = alice.patch(alice.org(f"/records/{detail['id']}"), json={"data": {"contact_email": "ceo@acme-security.test"}})
    assert ok.status_code == 200 and ok.json()["user_edited"]

    # Exclude Globex, export the rest.
    alice.post(alice.org(f"/jobs/{job['id']}/records/exclude"),
               json={"record_ids": [by_name["Globex Cyber Defense"]["id"]]})
    exp = alice.post(alice.org(f"/jobs/{job['id']}/exports"), json={}).json()
    exp = alice.get(alice.org(f"/exports/{exp['id']}")).json()
    assert exp["status"] == "ready", exp
    dl = alice.get(exp["download_url"])
    assert dl.status_code == 200
    assert "attachment" in dl.headers["content-disposition"] and dl.headers["cache-control"] == "no-store"
    wb = load_workbook(io.BytesIO(dl.content))
    data = wb["Data"]
    header = [c.value for c in data[1]]
    rows = [dict(zip(header, [c.value for c in row], strict=False)) for row in data.iter_rows(min_row=2)]
    labels = {f["name"]: f["label"] for f in schema["spec"]["fields"]}
    assert [r[labels["company_name"]] for r in rows] == ["Acme Security Ltd"]
    assert rows[0][labels["contact_email"]] == "ceo@acme-security.test"
    for ws in wb.worksheets:
        assert not any(c.data_type == "f" for row in ws.iter_rows() for c in row)

    usage = alice.get(alice.org("/usage")).json()
    assert usage["pages"] > 0 and usage["llm_requests"] == 0 and usage["export_rows"] == 1


def test_formula_payload_from_a_site_is_neutralized(make_user, crawl_env):
    bob = make_user("bob@example.com")
    project = bob.post(bob.org("/projects"), json={"name": "p"}).json()
    spec = {"entity": "company", "fields": [
        {"name": "company_name", "type": "string", "semantic": "company_name", "required": True},
        {"name": "description", "type": "text", "semantic": "description"}],
        "crawl": {"start_urls": [crawl_env["url"]("globex.test")], "max_pages": 5}}
    schema = bob.post(bob.org(f"/projects/{project['id']}/schemas"), json={"spec": spec}).json()
    job = bob.post(bob.org("/jobs"), json={"schema_id": schema["id"]}).json()
    recs = bob.get(bob.org(f"/jobs/{job['id']}/records")).json()["items"]
    assert recs[0]["data"]["description"].startswith("=HYPERLINK")  # stored as found
    exp = bob.post(bob.org(f"/jobs/{job['id']}/exports"), json={}).json()
    exp = bob.get(bob.org(f"/exports/{exp['id']}")).json()
    wb = load_workbook(io.BytesIO(bob.get(exp["download_url"]).content))
    values = [c.value for row in wb["Data"].iter_rows() for c in row]
    assert any(isinstance(v, str) and v.startswith("'=HYPERLINK") for v in values)
    assert not any(c.data_type == "f" for ws in wb.worksheets for row in ws.iter_rows() for c in row)


def test_cancel_and_retry(make_user, crawl_env, monkeypatch):
    from app.worker import tasks

    carol = make_user("carol@example.com")
    queued = []
    monkeypatch.setattr(tasks.run_job, "delay", lambda job_id: queued.append(job_id))
    project = carol.post(carol.org("/projects"), json={"name": "p"}).json()
    spec = {"fields": [{"name": "page_title", "type": "string", "semantic": "page_title"}],
            "crawl": {"start_urls": [crawl_env["url"]("acme-security.test")], "max_pages": 3}}
    schema = carol.post(carol.org(f"/projects/{project['id']}/schemas"), json={"spec": spec}).json()
    job = carol.post(carol.org("/jobs"), json={"schema_id": schema["id"]}).json()
    assert job["state"] == "created" and queued == [job["id"]]
    cancelled = carol.post(carol.org(f"/jobs/{job['id']}/cancel")).json()
    assert cancelled["state"] == "cancelled"
    assert carol.post(carol.org(f"/jobs/{job['id']}/cancel")).status_code == 409
    retried = carol.post(carol.org(f"/jobs/{job['id']}/retry")).json()
    assert retried["state"] == "queued"
    import uuid

    from app.services import pipeline
    from tests.conftest import test_resolver

    assert pipeline.run_job(uuid.UUID(job["id"]), resolver=test_resolver) in ("completed", "partially_completed")
