"""Score a recorded run's answer list against a reference list the user trusts.

For list questions ("which companies in X use Y") the useful measures are:
- recall: how many reference names the run found (with verified evidence);
- found but not in the reference: names to review by hand, since the reference may be incomplete;
- rejected candidates that match the reference: evidence the gate turned down, worth inspecting.

Usage:
    python -m argus.evals.list_benchmark exports/runs/<run> reference.xlsx [--column Company]

The reference is a CSV or XLSX supplied by the user; nothing about it is stored in the code.
"""

from __future__ import annotations

import argparse
import csv
import json
from pathlib import Path
import re
import sys
from typing import Any

from argus.investigation.claim_verification import is_verified

_SUFFIXES = (
    "incorporated", "inc", "corporation", "corp", "company", "co", "llc", "ltd", "limited", "plc",
    "lp", "llp", "gmbh", "ag", "sa", "nv", "bv", "pty", "group", "holdings", "the",
)


def normalize_name(name: str) -> str:
    """Comparable form of an organization name: case, punctuation and legal suffixes removed."""
    text = re.sub(r"\([^)]*\)", " ", (name or "").lower().replace("&amp;", "&"))  # "(Atlanta office)"
    words = re.sub(r"[^a-z0-9& ]+", " ", text).split()
    while words and words[-1] in _SUFFIXES:
        words.pop()
    while words and words[0] == "the":
        words.pop(0)
    return " ".join(words)


_NAME_HEADERS = ("company", "company name", "organization", "organisation", "organization name", "name")


def load_reference(path: Path, column: str | None = None) -> list[str]:
    """Names from a reference table: in every sheet, the first row (of the top 15) with a
    company/name header cell (or the given column header), then the non-empty cells below it.
    Names are merged across sheets in order, without duplicates."""
    tables: list[list[list[Any]]]
    if path.suffix.lower() in (".xlsx", ".xlsm"):
        from openpyxl import load_workbook

        wb = load_workbook(path, read_only=True, data_only=True)
        tables = [[list(r) for r in ws.iter_rows(values_only=True)] for ws in wb.worksheets]
    else:
        with path.open(newline="", encoding="utf-8-sig") as fh:
            tables = [list(csv.reader(fh))]
    wanted = (column.strip().lower(),) if column else _NAME_HEADERS
    names: list[str] = []
    seen: set[str] = set()
    for rows in tables:
        for i, row in enumerate(rows[:15]):
            cells = [str(c or "").strip().lower() for c in row]
            col = next((j for j, c in enumerate(cells) if c in wanted), None)
            if col is None:
                continue
            for r in rows[i + 1:]:
                value = str(r[col]).strip() if col < len(r) and r[col] is not None else ""
                if value and normalize_name(value) not in seen:
                    seen.add(normalize_name(value))
                    names.append(value)
            break
    if not names:
        raise ValueError(f"No column {'named ' + repr(column) if column else 'with a company/name header'} in {path}")
    return names


def answers_from_result(result: dict[str, Any], predicate_prefix: str = "USES") -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    """(verified answer claims, rejected answer claims) of a run's result."""
    claims = [c for c in result.get("claims", []) if str(c.get("predicate", "")).upper().startswith(predicate_prefix.upper())]
    return [c for c in claims if is_verified(c)], [c for c in claims if not is_verified(c)]


def _same_org(a: str, b: str) -> bool:
    """Equal normalized names, or one is the other's leading words ("lockheed martin" /
    "lockheed martin aeronautics") when the shorter is specific enough to identify it."""
    if a == b:
        return True
    short, long_ = sorted((a, b), key=len)
    specific = len(short.split()) >= 2 or len(short) >= 6
    return specific and (long_ + " ").startswith(short + " ")


def _match(name: str, ref: dict[str, str]) -> str | None:
    key = normalize_name(name)
    if not key:
        return None
    if key in ref:
        return ref[key]
    return next((orig for k, orig in ref.items() if _same_org(key, k)), None)


def score(result: dict[str, Any], reference: list[str], predicate_prefix: str = "USES") -> dict[str, Any]:
    verified, rejected = answers_from_result(result, predicate_prefix)
    ref = {normalize_name(n): n for n in reference if normalize_name(n)}
    names = sorted({c.get("subject", "") for c in verified if normalize_name(c.get("subject", ""))})
    matched = {n: _match(n, ref) for n in names}
    hits = sorted({m for m in matched.values() if m})
    extra = [n for n, m in matched.items() if not m]
    rejected_hits = {m for c in rejected if (m := _match(c.get("subject", ""), ref))}
    return {
        "objective": result.get("objective"),
        "reference_size": len(ref),
        "verified_answers": len(names),
        "recall": round(len(hits) / len(ref), 3) if ref else None,
        "found_in_reference": hits,
        "found_not_in_reference": extra,
        "rejected_but_in_reference": sorted(rejected_hits - set(hits)),
    }


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument("run_dir", help="Recorded run folder (contains result.json) or a result.json file")
    parser.add_argument("reference", help="Reference list (.xlsx or .csv)")
    parser.add_argument("--column", help="Reference column holding the names")
    parser.add_argument("--predicate", default="USES", help="Answer claims' predicate prefix (default USES)")
    args = parser.parse_args(argv)
    path = Path(args.run_dir)
    result = json.loads((path / "result.json" if path.is_dir() else path).read_text(encoding="utf-8"))
    report = score(result, load_reference(Path(args.reference), args.column), args.predicate)
    json.dump(report, sys.stdout, indent=2)
    sys.stdout.write("\n")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
