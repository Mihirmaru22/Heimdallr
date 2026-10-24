"""Companies whose own SEC filings show they use a product, with every piece of proof kept or rejected."""

import json
from datetime import date, timedelta

import httpx
from openpyxl import load_workbook
from sqlalchemy import select

from app.db import session_scope
from app.discovery.sec_filings import SecFilingsProvider, company_name, document_url, website_in
from app.evidence.judge import Judge
from app.llm.openrouter import Reply
from app.models import Evidence
from app.services.diagnostics import run_diagnostics
from app.spec.models import ExtractionSpec

RECENT = (date.today() - timedelta(days=200)).isoformat()
OLDER = (date.today() - timedelta(days=400)).isoformat()


def hit(cik, name, adsh, filename, form, day, states=("OH",), file_type=None):
    return {"_id": f"{adsh}:{filename}", "_source": {
        "ciks": [f"{cik:010d}"], "display_names": [f"{name} (TCK) (CIK {cik:010d})"], "form": form,
        "file_type": file_type or form, "file_date": day, "biz_states": list(states),
        "biz_locations": ["Avon Lake, OH"], "sics": ["2821"], "adsh": adsh}}


HITS = [
    hit(1122976, "AVIENT CORP", "0001122976-25-000010", "avnt-10k.htm", "10-K", RECENT),
    hit(1122976, "AVIENT CORP", "0001122976-24-000020", "avnt-10q.htm", "10-Q", OLDER),  # a 10-K beats it
    hit(1420720, "Rimini Street, Inc.", "0001420720-25-000003", "rmni-10k.htm", "10-K", RECENT, states=("NV",)),
    hit(1057379, "HACKETT GROUP, INC.", "0001057379-25-000004", "hckt-10k.htm", "10-K", RECENT, states=("FL",)),
    hit(1300000, "MAPLE MILLS LTD", "0001300000-25-000001", "mml-20f.htm", "20-F", RECENT, states=("A6",)),
    hit(1400000, "QUIET CO", "0001400000-25-000001", "qc-10k.htm", "10-K", RECENT, states=("TX",)),
    hit(1500000, "GLOBEX INC", "0001500000-25-000001", "gbx-10k.htm", "10-K", RECENT, states=("CA",)),
    hit(1438569, "Grifols SA", "0001438569-26-000001", "grf-20f.htm", "20-F", RECENT, states=("NY",)),
]
DOCS = {
    "avnt-10k.htm": "<html><body><p>Avient makes specialty polymers. In 2025 we completed the migration of our "
                    "plants to SAP S/4HANA, which now runs order to cash. Revenue grew.</p></body></html>",
    "rmni-10k.htm": "<p>We compete with SAP for support of SAP ECC and S/4HANA customers.</p>",
    "hckt-10k.htm": "<p>We offer implementation services for SAP S/4HANA to large enterprises.</p>",
    "qc-10k.htm": "<p>Nothing here names the product, only the search index thought so.</p>",
    "gbx-10k.htm": "<p>Our systems include SAP S/4HANA.</p>",
}


def sec_and_wikidata(port, seen):
    def handler(request: httpx.Request) -> httpx.Response:
        seen.append(str(request.url))
        if request.url.host == "efts.sec.gov":
            q = request.url.params["q"]
            if q.count('"') > 2 and " OR " not in q:  # like EDGAR: names separated by a space must all appear
                return httpx.Response(200, json={"hits": {"total": {"value": 0}, "hits": []}})
            start = int(request.url.params.get("from", 0))
            page = HITS[start:start + 4]
            return httpx.Response(200, json={"hits": {"total": {"value": len(HITS)}, "hits": page}})
        if request.url.host == "www.sec.gov":
            name = request.url.path.rsplit("/", 1)[1]
            return httpx.Response(200, text=DOCS[name]) if name in DOCS else httpx.Response(404)
        if request.url.params.get("list") == "search":  # Wikidata: which items carry these CIKs
            qids = [q for cik, q in (("0001122976", "Q1"), ("0001500000", "Q2"))
                    if f"P5531={cik}" in request.url.params["srsearch"]]
            return httpx.Response(200, json={"query": {"search": [{"title": q} for q in qids]}})
        ents = {"Q1": ("Avient", "0001122976", f"http://acme-security.test:{port}/"),
                "Q2": ("Globex", "0001500000", f"http://globex.test:{port}/")}
        return httpx.Response(200, json={"entities": {q: {
            "labels": {"en": {"value": label}}, "claims": {
                "P5531": [{"mainsnak": {"datavalue": {"value": cik}}}],
                "P856": [{"mainsnak": {"datavalue": {"value": site}}}]}}
            for q, (label, cik, site) in ents.items() if q in request.url.params["ids"].split("|")}})
    return handler


class Router:
    """Avient: uses it. Globex: an answer that cannot be read, so its proof stays unreviewed."""

    def __init__(self):
        self.calls = []

    async def chat(self, messages, **_kw):
        user = messages[-1]["content"]
        self.calls.append(user)
        if "GLOBEX" in user or "Globex" in user:
            return Reply("sorry, I cannot help", "minimax/minimax-m3.1-flash:free")
        return Reply(json.dumps({"relationship": "uses", "quote": 1, "reason": "runs its plants on it"}),
                     "minimax/minimax-m3.1-flash:free")

    async def aclose(self):
        pass


def spec(**discovery):
    return ExtractionSpec.model_validate({"entity": "company", "fields": [
        {"name": "company_name", "type": "string", "semantic": "company_name", "required": True},
        {"name": "website_url", "type": "url", "semantic": "website_url"},
        {"name": "uses_technology", "type": "string", "required": True},
        {"name": "evidence", "type": "text"}, {"name": "evidence_url", "type": "url"},
        {"name": "evidence_date", "type": "date"}, {"name": "evidence_type", "type": "string"}],
        "discovery": {"providers": ["sec_filings"], "technology": "SAP", "countries": ["US"], **discovery}})


def test_the_website_a_filing_names_for_its_own_company():
    release = ("Celanese is moving to SAP S/4HANA. Source: Business Wire (www.businesswire.com). See sap.com. "
               "About Celanese: visit www.celanese.com or investors.celanese.com. Media: media@celanese.com")
    assert website_in(release, "Celanese Corp") == "https://www.celanese.com/"
    assert website_in("Our Internet address is www.agcocorp.com.", "AGCO CORP /DE") == "https://www.agcocorp.com/"
    assert website_in("Visit https://www.hbfuller.com today.", "H.B. Fuller Co") == "https://www.hbfuller.com/"
    assert website_in("See www.spectrumbrands.com.", "Spectrum Brands Holdings, Inc.") == "https://www.spectrumbrands.com/"
    assert website_in("Our partner www.sap.com and www.sec.gov.", "Coty Inc.") == ""  # never another company's
    assert website_in("See investors.coty.com.", "Coty Inc.") == "https://www.coty.com/"  # not the IR host
    assert website_in("Visit gov2x.com.", "V2X, Inc.") == "https://gov2x.com/"  # the host as written


async def test_the_filing_s_own_website_beats_older_data_and_spares_the_wikidata_lookup(db_clean):
    from app.models import SecCompany

    with session_scope(bypass=True) as db:  # SEC's company data still names the old parent's parked domain
        db.add(SecCompany(cik=109177, name="Spectrum Brands Holdings, Inc.", rank=1, website="https://hrggroup.com/"))
    asked = []
    client = httpx.AsyncClient(transport=httpx.MockTransport(lambda r: asked.append(r) or httpx.Response(500)))
    users = {"0000109177": {"company": "Spectrum Brands", "named_website": "https://www.spectrumbrands.com/"}}
    await SecFilingsProvider(client, Judge(Router()), gap=0)._attach_websites(client, users)
    assert users["0000109177"]["website"] == "https://www.spectrumbrands.com/" and not asked


def test_display_names_and_document_links():
    assert company_name("AVIENT CORP (AVNT) (CIK 0001122976)") == "AVIENT CORP"
    assert document_url("0001122976", "0001122976-25-000010:avnt-10k.htm") == \
        "https://www.sec.gov/Archives/edgar/data/1122976/000112297625000010/avnt-10k.htm"


async def test_only_users_become_results_and_every_rejection_says_why(db_clean):
    seen = []
    client = httpx.AsyncClient(transport=httpx.MockTransport(sec_and_wikidata(8080, seen)))
    router = Router()
    provider = SecFilingsProvider(client, Judge(router), gap=0)
    results = await provider.discover(spec())
    await client.aclose()

    search = [u for u in seen if "efts.sec.gov" in u]
    assert len(search) == 2  # paged
    q = httpx.URL(search[0]).params["q"]
    assert q.startswith('"SAP S/4HANA" OR "S/4HANA" OR ') and '"SAP"' not in q  # any quoted name, never "SAP" alone
    (avient, globex) = sorted(results, key=lambda r: r.url)
    assert avient.url == "http://acme-security.test:8080/" and avient.method == "sec_filings"
    assert avient.hint["company_name"] == "Avient" and avient.hint["uses_technology"] == "Confirmed"
    assert avient.hint["evidence"].startswith("In 2025 we completed the migration of our plants to SAP S/4HANA")
    assert avient.hint["evidence_url"].endswith("/1122976/000112297625000010/avnt-10k.htm")
    assert avient.hint["evidence_date"] == RECENT and avient.hint["evidence_type"] == "SEC filing (10-K)"
    assert globex.hint["uses_technology"] == "Mentioned"  # unreviewed proof: never confirmed

    why = {e["company"]: e["reason"] for e in provider.evidence if not e["kept"]}
    assert why["Rimini Street, Inc."].startswith("competitor")
    assert why["HACKETT GROUP, INC."].startswith("partner or reseller")
    assert why["MAPLE MILLS LTD"] == "business address outside the US"
    assert why["Grifols SA"] == "business address outside the US"  # a 20-F filer, whatever its agent's address
    assert why["QUIET CO"] == "the product's name was not found in the document text"
    # Rules settled Avient (plain own use), the competitor and the reseller; Globex's unreadable answer was
    # asked of another model once.
    assert len(router.calls) == 2
    assert not any("avnt-10q" in u for u in seen)  # one document per company: its 10-K


async def test_sec_flaking_is_tried_again_and_a_lasting_failure_is_not_reported_as_zero_filings(db_clean):
    from app.discovery.providers import DiscoveryError

    calls = []
    good = sec_and_wikidata(8080, [])

    def flaky(fails):
        def handler(request: httpx.Request) -> httpx.Response:
            if request.url.host == "efts.sec.gov":
                calls.append(1)
                if len(calls) == 1:
                    raise httpx.ConnectTimeout("timed out")
                if len(calls) <= fails:
                    return httpx.Response(500, json={"message": "Internal server error"})
            return good(request)
        return handler

    provider = SecFilingsProvider(httpx.AsyncClient(transport=httpx.MockTransport(flaky(2))), Judge(Router()), gap=0)
    lines = []
    provider.on_note = lambda kind, message, url=None, **_: lines.append((kind, message))
    assert len(await provider.discover(spec())) == 2  # a timeout and a 500, then SEC answers
    assert ("failed", "SEC did not answer (efts.sec.gov); trying again in 0s") in lines

    calls.clear()
    provider = SecFilingsProvider(httpx.AsyncClient(transport=httpx.MockTransport(flaky(99))), Judge(Router()), gap=0)
    import pytest
    with pytest.raises(DiscoveryError, match="SEC full-text search failed"):
        await provider.discover(spec())
    assert len(calls) == 4  # the first try and three more


async def test_outside_the_us_it_says_so(db_clean):
    import pytest

    from app.discovery.providers import DiscoveryError

    with pytest.raises(DiscoveryError, match="listed in the US"):
        await SecFilingsProvider(httpx.AsyncClient(), Judge(Router()), gap=0).discover(spec(countries=["DE"]))


def test_a_technology_job_keeps_proven_users_and_exports_the_proof(db_clean, crawl_env, monkeypatch, tmp_path,
                                                                    capsys):
    from app.crawl import runner
    from app.export.xlsx import build_workbook
    from app.services import exports, pipeline
    from tests.integration.test_pipeline import activity, make_job, records

    port = crawl_env["port"]
    client = httpx.AsyncClient(transport=httpx.MockTransport(sec_and_wikidata(port, [])))
    monkeypatch.setattr(runner, "providers_for",
                        lambda s: [SecFilingsProvider(client, Judge(Router()), gap=0)])
    s = spec().model_dump(mode="json")
    job_id, org_id = make_job([], fields=s["fields"], discovery=s["discovery"])
    assert pipeline.run_job(job_id, resolver=crawl_env["resolver"]) == "completed"

    by_site = {r.data["website_url"]: r for r in records(job_id, org_id)}
    avient = by_site[f"http://acme-security.test:{port}/"]
    assert avient.data["uses_technology"] == "Confirmed" and not avient.excluded
    assert avient.data["evidence"].startswith("In 2025 we completed")
    assert by_site[f"http://globex.test:{port}/"].excluded  # only a mention: shown, but left out of exports

    lines = activity(job_id, org_id)
    assert any(k == "search" and "SEC full-text search: 8 filing(s) from 7 companies" in m for k, m in lines)
    assert any(k == "read" and m.startswith("Reading AVIENT CORP's 10-K") for k, m in lines)
    assert any(k == "kept" and m.startswith("Proof: AVIENT CORP's 10-K") and "(Confirmed)" in m for k, m in lines)
    assert any(k == "dropped" and m.startswith("Rejected Rimini Street, Inc.'s 10-K: competitor") for k, m in lines)

    with session_scope(org_id) as db:
        saved = db.execute(select(Evidence).where(Evidence.job_id == job_id)).scalars().all()
        assert len(saved) == 7 and sum(e.kept for e in saved) == 2
        proof = exports._evidence(db, job_id)
        diagnostics = run_diagnostics(db, job_id)
    path = build_workbook(tmp_path / "x.xlsx", ExtractionSpec.model_validate(s), [], {"Job": "t"},
                          evidence=proof, diagnostics=diagnostics).path
    wb = load_workbook(path)
    assert wb.sheetnames == ["Data", "Sources", "Evidence", "Rejected", "Run Summary", "Diagnostics"]
    diag = [r for r in wb["Diagnostics"].iter_rows(min_row=2, values_only=True)]
    assert ("Plan", "providers", "sec_filings") in diag and ("Plan", "industry", "(none)") in diag
    assert any(sec == "Source: sec_filings" and v.startswith("SEC full-text search: 8 filing(s)")
               for sec, _i, v in diag)
    ev = list(wb["Evidence"].iter_rows(values_only=True))
    assert ev[0][:4] == ("Company", "Website", "Product", "Verdict")
    assert ev[1][0] == "Avient" and ev[1][3] == "Confirmed"
    rejected = {r[0]: r[2] for r in wb["Rejected"].iter_rows(min_row=2, values_only=True)}
    assert rejected["MAPLE MILLS LTD"] == "business address outside the US"

    from app.services.scorecard import score_job
    key = {"must_include": [{"company": "Avient", "website": f"http://acme-security.test:{port}/"},
                            {"company": "Lancaster Colony"}, {"company": "Hackett Group"}],
           "must_exclude": [{"company": "Rimini Street"}, {"company": "Globex"}]}
    with session_scope(org_id) as db:
        card = score_job(db, job_id, key)
    assert card["proof: companies kept"] == "2" and card["proof lost: competitor"] == "1"
    assert card["proof lost: outside the requested country"] == "2"
    assert card["records included"] == "1" and card["quote names the product"] == "1/1 (100%)"
    assert card["key: must-have companies found"] == "1/3 (33%)"
    assert card["key: missing"] == "Lancaster Colony, Hackett Group"
    assert card["key: lost at: proof rejected: partner or reseller"] == "1 (Hackett Group)"
    assert card["key: rows that must not be there"] == "0"  # Globex is shown but excluded
    assert card["key: website correct"] == "1/1 (100%)"
    assert card["key: lost at: not found by any source"] == "1 (Lancaster Colony)"

    from app.cli import job_report
    job_report(str(job_id))
    out = capsys.readouterr().out
    assert "providers                        sec_filings\n" in out and "SEC full-text search: 8 filing(s)" in out
    assert "SEC filing (10-K) kept" in out and "== Pages per site" in out


async def test_a_company_with_proof_and_no_known_website_is_looked_up_by_name(db_clean, monkeypatch):
    from app.config import get_settings

    monkeypatch.setattr(get_settings(), "searxng_url", "http://searx")

    def answer(request):
        if request.url.host == "searx":
            return httpx.Response(200, json={"results": [{"url": "https://www.accuray.com/"}] * 5})
        return httpx.Response(200, json={"query": {"search": []}})  # Wikidata has no item for the CIK

    async def fetch_text(url):
        return "Accuray radiation therapy" if url == "https://www.accuray.com/" else None

    client = httpx.AsyncClient(transport=httpx.MockTransport(answer))
    provider = SecFilingsProvider(client, Judge(Router()), gap=0)
    provider.fetch_text = fetch_text
    users = {"0001138723": {"company": "ACCURAY INC"}}
    await provider._attach_websites(client, users)
    assert users["0001138723"]["website"] == "https://www.accuray.com/"
