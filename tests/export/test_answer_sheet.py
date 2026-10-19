"""The Answers sheet: one row per thing found, its evidence beside it, columns to work in."""

from openpyxl import load_workbook
import pytest

from argus.export.answer_rows import build_answer_rows, evidence_strength, evidence_type_of
from argus.export.excel_exporter import OSINTExcelExporter

LIST_ENUM = {
    "raw": "Which hospitals in Ohio use Epic?", "answer_type": "ORGANIZATION", "answer_noun": "hospitals",
    "relation": "uses", "relation_phrase": "use", "object": "Epic", "geography": "Ohio",
}


def claim(subject, predicate, obj, status="SUPPORTED", url="", quote="", eid="e1"):
    return {"claim_id": f"{subject}-{predicate}", "subject": subject, "predicate": predicate, "object": obj,
            "status": status, "citation": url, "evidence_id": eid, "source_passage": quote}


def list_run(**extra):
    data = {
        "objective": "Which hospitals in Ohio use Epic?", "enumeration_objective": LIST_ENUM,
        "stop_reason": "OBJECTIVE_SATISFIED",
        "claims": [
            claim("Mercy General", "uses_software", "Epic", "CORROBORATED",
                  "https://jobs.lever.co/mercy/epic-analyst", "Epic Ambulatory analyst, Columbus.", "e1"),
            claim("Mercy General", "uses_software", "Epic", "SUPPORTED",
                  "https://mercygeneral.example/it", "Our Epic rollout finished in 2024.", "e2"),
            claim("Mercy General", "website", "https://mercygeneral.example/", "SUPPORTED", "", "", "e2"),
            claim("Lakeside Clinic", "uses_software", "Epic", "SUPPORTED",
                  "https://lakeside.example/careers", "Epic MyChart support specialist.", "e3"),
            claim("Lakeside Clinic", "headquartered_in", "Toledo, OH", "SUPPORTED", "", "", "e3"),
        ],
        "evidence": [{"evidence_id": "e1", "url": "https://jobs.lever.co/mercy/epic-analyst", "title": "Epic analyst"},
                     {"evidence_id": "e2", "url": "https://mercygeneral.example/it", "title": "IT"},
                     {"evidence_id": "e3", "url": "https://lakeside.example/careers", "title": "Careers"}],
        "entities": {}, "research_actions": [{"query": "Ohio hospital Epic analyst"}], "audit_trail": [],
        "relationships": [], "unresolved_questions": [], "limitations": ["Indeed and Glassdoor were not read."],
    }
    data.update(extra)
    return data


def sheet(tmp_path, data, name="Answers"):
    path = OSINTExcelExporter(export_dir=tmp_path).export_investigation(data)
    return load_workbook(path)[name]


def test_one_row_per_answer_with_its_strongest_quote():
    columns, rows = build_answer_rows(list_run())
    assert [r.answer for r in rows] == ["Mercy General", "Lakeside Clinic"]  # strongest first
    mercy = rows[0]
    assert (mercy.evidence_strength, mercy.sources) == ("Strong", 2)
    assert mercy.evidence_says == "Epic Ambulatory analyst, Columbus."
    assert mercy.evidence_url == "https://jobs.lever.co/mercy/epic-analyst"
    assert "https://mercygeneral.example/it" in mercy.other_sources()
    assert rows[1].evidence_strength == "Moderate"
    assert columns == ["Website", "Where"]


def test_attributes_fill_columns_instead_of_becoming_their_own_rows():
    _, rows = build_answer_rows(list_run())
    assert rows[0].attributes["Website"] == "https://mercygeneral.example/"
    assert rows[1].attributes["Where"] == "Toledo, OH"
    assert all("mercygeneral.example/" != r.answer for r in rows)


def test_attribute_columns_are_absent_when_nothing_found_them():
    data = list_run(claims=[claim("Mercy General", "uses_software", "Epic", "SUPPORTED",
                                  "https://x.example/a", "Epic analyst.", "e1")])
    columns, rows = build_answer_rows(data)
    assert columns == [] and len(rows) == 1


def test_a_question_that_is_not_a_list_answers_with_the_object():
    data = list_run(objective="Who is the chief executive of Acme?", enumeration_objective=None,
                    claims=[claim("Acme", "chief_executive_officer", "Jane Doe", "SUPPORTED",
                                  "https://acme.example/board", "Jane Doe is chief executive.", "e1")])
    _, rows = build_answer_rows(data)
    assert rows[0].answer == "Jane Doe" and rows[0].detail == "Acme chief_executive_officer"


def test_gate_rejected_candidates_never_reach_the_answer_rows():
    data = list_run()
    data["claims"].append(claim("Nope Ltd", "uses_software", "Epic", "GATE_REJECTED", "https://n.example", "no", "e9"))
    _, rows = build_answer_rows(data)
    assert "Nope Ltd" not in [r.answer for r in rows]


@pytest.mark.parametrize("url,expected", [
    ("https://jobs.lever.co/x/y", "Job post"),
    ("https://boards.greenhouse.io/x", "Job post"),
    ("https://www.sec.gov/edgar/filing/123", "Filing"),
    ("https://siemens.com/case-studies/acme", "Case study"),
    ("https://web.archive.org/web/2020/https://x.example", "Archived copy"),
    ("https://en.wikipedia.org/wiki/Epic", "Reference page"),
    ("https://acme.example/about", "Web page"),
])
def test_evidence_type_reads_the_url(url, expected):
    assert evidence_type_of(url) == expected


@pytest.mark.parametrize("status,sources,quote,expected", [
    ("SUPPORTED", 2, True, "Strong"),
    ("CORROBORATED", 1, True, "Strong"),
    ("SUPPORTED", 1, True, "Moderate"),
    ("SUPPORTED", 1, False, "Weak"),
])
def test_evidence_strength(status, sources, quote, expected):
    assert evidence_strength(status, sources, quote) == expected


def test_the_sheet_is_laid_out_to_be_read_and_worked_in(tmp_path):
    ws = sheet(tmp_path, list_run())
    assert ws.cell(row=1, column=1).value == "Which hospitals in Ohio use Epic?"
    assert "2 hospitals" in ws.cell(row=2, column=1).value
    headers = [c.value for c in ws[4]]
    assert headers[0] == "Hospitals" and headers[-3:] == ["Status", "Owner", "Notes"]
    assert ws.freeze_panes == "A5" and ws.auto_filter.ref.startswith("A4:")
    assert ws.cell(row=4, column=1).fill.fgColor.rgb.endswith("1F3864")  # the navy header band
    url_col = headers.index("Evidence URL") + 1
    assert ws.cell(row=5, column=url_col).hyperlink is not None
    status_col = headers.index("Status") + 1
    assert ws.cell(row=5, column=status_col).fill.fill_type == "solid"  # shaded: yours to fill in


def test_answers_dashboard_counts_the_sheet_and_method_sheet_says_how(tmp_path):
    path = OSINTExcelExporter(export_dir=tmp_path).export_investigation(list_run())
    wb = load_workbook(path)
    assert wb.sheetnames[:3] == ["Answers", "Dashboard", "Sources & Method"]
    dash = {r[0].value: r[1].value for r in wb["Dashboard"].iter_rows(min_row=5, max_row=11)}
    assert dash["Answers found"] == "=COUNTA(Answers!A5:A6)"
    assert dash["Strong evidence (two or more sites)"] == '=COUNTIF(Answers!$E$5:$E$6,"Strong")'
    method = "\n".join(str(r[0].value) for r in wb["Sources & Method"].iter_rows(min_row=1, max_row=30))
    assert "Ohio hospital Epic analyst" in method
    assert "Indeed and Glassdoor were not read." in method
    assert "lakeside.example" in method


def test_a_run_that_found_nothing_says_so_instead_of_showing_an_empty_table(tmp_path):
    ws = sheet(tmp_path, list_run(claims=[], evidence=[]))
    assert "Nothing passed the evidence check" in ws.cell(row=5, column=1).value


def test_the_markdown_report_opens_with_the_same_answer_table():
    from argus.export.markdown_exporter import OSINTMarkdownExporter

    md = OSINTMarkdownExporter(export_dir="/tmp")._render_markdown_report(list_run(), "Epic")
    assert "## Answers (2 hospitals)" in md
    assert "| Hospitals | Website | Where | Evidence | Strength | Sources |" in md
    assert "[Job post](https://jobs.lever.co/mercy/epic-analyst)" in md
    assert md.index("## Answers") < md.index("## Key Corroborated Findings")


def test_a_cdn_host_picked_up_in_passing_is_not_an_answer():
    data = list_run(claims=[
        claim("Mercy General", "uses_software", "Epic", "SUPPORTED", "https://x.example/a", "Epic analyst.", "e1"),
        claim("rr1---sn-4g5ednss.googlevideo.com", "serves_content_for", "Epic", "SUPPORTED",
              "https://x.example/a", "video host", "e1"),
    ])
    _, rows = build_answer_rows(data)
    assert [r.answer for r in rows] == ["Mercy General"]
